"""
学习/项目进度智能总结系统 V1.2
作者：[林怡佳]
开发日期：2026-07-01
更新日期：2026-09-18

系统架构：多Agent协作流水线 + 跨周期诊断
- CollectorAgent（采集Agent）：整理和分类用户的零散输入记录
- AnalyzerAgent（分析Agent）：深度分析单期进展、问题与规律
- FormatterAgent（格式化Agent）：生成结构化报告并支持导出Word文档
- DiagnosticAgent（诊断Agent，V1.2新增）：跨周期识别重复问题模式，
  给出有据可查的改进建议（模式识别 + 诊断建议 两阶段）

V1.2 新增功能：
- 跨周期智能诊断：分析最近N期历史报告，发现反复出现的问题并溯源
- 诊断结果持久化存储（data/diagnoses.json）
- /api/diagnose 接口

V1.1 功能：
- 历史报告持久化存储（JSON文件，程序重启不丢失）
- 历史报告列表查询接口（支持分页）
- 历史报告详情查询接口
- 历史报告删除接口
- 统计数据接口（总报告数、各周期分布）

技术栈：
- 后端框架：FastAPI
- AI接口：DeepSeek API（兼容OpenAI SDK）
- 文档生成：python-docx
- 数据持久化：JSON文件存储
- 前端：HTML5 + Vanilla JavaScript

本文件只负责：应用初始化、依赖装配（Agent与客户端）、API路由。
具体的数据模型见 models.py，存储逻辑见 storage.py，Agent实现见 agents.py。
"""

import io
import logging
import os
from datetime import datetime

from dotenv import load_dotenv
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from fastapi import FastAPI, HTTPException, Query
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openai import OpenAI

from agents import CollectorAgent, AnalyzerAgent, FormatterAgent, DiagnosticAgent
from models import (
    SummarizeRequest, SummarizeResponse,
    ExportRequest,
    ReportListResponse, ReportSummary,
    StatsResponse,
    DiagnoseRequest, DiagnoseResponse, ProblemPattern, DiagnosticSuggestion,
)
from storage import (
    init_storage,
    load_reports, save_report, delete_report_by_id, get_report_by_id,
    get_recent_reports, get_reports_by_ids,
    save_diagnosis,
)

# ============================================================
# 初始化配置
# ============================================================

load_dotenv()

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s"
)
logger = logging.getLogger("progress_agent")

DEEPSEEK_API_KEY = os.environ.get("DEEPSEEK_API_KEY", "")

deepseek_client = OpenAI(
    api_key=DEEPSEEK_API_KEY,
    base_url="https://api.deepseek.com"
)

app = FastAPI(
    title="个人进度智能总结系统",
    description="基于多Agent协作架构的学习/项目进度追踪、报告生成与跨周期诊断系统",
    version="1.2.0"
)

# 启动时初始化存储（reports.json + diagnoses.json）
init_storage()

# 内存缓存（用于导出Word时快速访问当次生成的报告）
report_cache: dict = {}


# ============================================================
# Agent 协调器（Pipeline 控制器）
# ============================================================

class AgentPipeline:
    """
    多Agent协作流水线控制器（单周期总结）。
    负责按顺序调度 Collector -> Analyzer -> Formatter 三个Agent。
    """

    def __init__(self, client: OpenAI):
        self.collector = CollectorAgent(client)
        self.analyzer = AnalyzerAgent(client)
        self.formatter = FormatterAgent(client)
        self.logger = logging.getLogger("AgentPipeline")

    def run(self, raw_notes: str, period: str, author_name: str = "") -> SummarizeResponse:
        """执行完整的多Agent协作流程，返回三个Agent结果和最终报告"""
        self.logger.info(f"启动多Agent流水线，周期：{period}")

        self.logger.info("=== 阶段1：CollectorAgent 运行中 ===")
        collector_result = self.collector.run(raw_notes)
        if collector_result.status == "error":
            raise HTTPException(status_code=500, detail=f"采集阶段失败: {collector_result.output}")

        self.logger.info("=== 阶段2：AnalyzerAgent 运行中 ===")
        analyzer_result = self.analyzer.run(collector_result.output)
        if analyzer_result.status == "error":
            raise HTTPException(status_code=500, detail=f"分析阶段失败: {analyzer_result.output}")

        self.logger.info("=== 阶段3：FormatterAgent 运行中 ===")
        formatter_result = self.formatter.run(
            collected_json=collector_result.output,
            analysis_text=analyzer_result.output,
            period=period
        )
        if formatter_result.status == "error":
            raise HTTPException(status_code=500, detail=f"格式化阶段失败: {formatter_result.output}")

        report_id = datetime.now().strftime("%Y%m%d_%H%M%S")

        report_data = {
            "report_id": report_id,
            "period": period,
            "author_name": author_name,
            "created_at": datetime.now().isoformat(),
            "final_report": formatter_result.output,
            "collector_output": collector_result.output,
            "analyzer_output": analyzer_result.output,
            "raw_notes": raw_notes
        }

        save_report(report_data)
        report_cache[report_id] = report_data

        self.logger.info(f"流水线执行完成，报告ID：{report_id}")

        return SummarizeResponse(
            collector_result=collector_result,
            analyzer_result=analyzer_result,
            formatter_result=formatter_result,
            final_report=formatter_result.output,
            report_id=report_id
        )


class DiagnosticPipeline:
    """
    诊断流水线控制器（跨周期诊断，V1.2新增）。
    负责取出最近N条历史报告，交给 DiagnosticAgent 做两阶段分析，
    并将诊断结果持久化。
    """

    def __init__(self, client: OpenAI):
        self.diagnostic_agent = DiagnosticAgent(client)
        self.logger = logging.getLogger("DiagnosticPipeline")

    def run(self, report_ids: list = None, report_count: int = 4, author_name: str = None) -> DiagnoseResponse:
        if report_ids:
            self.logger.info(f"启动诊断流水线，分析用户手动勾选的 {len(report_ids)} 条报告")
            recent_reports = get_reports_by_ids(report_ids)
        else:
            self.logger.info(f"启动诊断流水线，分析最近 {report_count} 条报告")
            recent_reports = get_recent_reports(count=report_count, author_name=author_name)

        if len(recent_reports) < 2:
            raise HTTPException(
                status_code=400,
                detail=f"可用于诊断的报告不足2条（当前{len(recent_reports)}条），"
                       f"请检查勾选的报告ID是否正确，或先积累更多历史报告"
            )

        self.logger.info("=== DiagnosticAgent 运行中（模式识别 -> 诊断建议）===")
        diagnostic_result = self.diagnostic_agent.run(recent_reports)

        if diagnostic_result.status == "error":
            raise HTTPException(status_code=500, detail=f"诊断阶段失败: {diagnostic_result.output}")

        import json
        result_data = json.loads(diagnostic_result.output)

        diagnosis_id = datetime.now().strftime("diag_%Y%m%d_%H%M%S")

        patterns = [ProblemPattern(**p) for p in result_data.get("patterns", [])]
        suggestions = [DiagnosticSuggestion(**s) for s in result_data.get("suggestions", [])]

        response = DiagnoseResponse(
            diagnosis_id=diagnosis_id,
            analyzed_report_count=len(recent_reports),
            analyzed_report_ids=[r.get("report_id") for r in recent_reports],
            patterns=patterns,
            suggestions=suggestions,
            overall_summary=result_data.get("overall_summary", ""),
            created_at=datetime.now().isoformat()
        )

        save_diagnosis(response.model_dump())

        self.logger.info(f"诊断流水线执行完成，诊断ID：{diagnosis_id}")

        return response


pipeline = AgentPipeline(deepseek_client)
diagnostic_pipeline = DiagnosticPipeline(deepseek_client)


# ============================================================
# API 路由
# ============================================================

@app.get("/")
def index():
    """返回前端首页"""
    return FileResponse("static/index.html")


@app.get("/health")
def health_check():
    """健康检查接口"""
    return {
        "status": "ok",
        "version": "1.2.0",
        "api_key_configured": bool(DEEPSEEK_API_KEY),
        "total_reports": len(load_reports())
    }


@app.post("/api/summarize", response_model=SummarizeResponse)
def summarize(req: SummarizeRequest):
    """核心接口：触发多Agent流水线生成进度报告"""
    if not DEEPSEEK_API_KEY:
        raise HTTPException(status_code=400, detail="未配置 API Key，请检查 .env 文件")
    if not req.raw_notes.strip():
        raise HTTPException(status_code=400, detail="输入内容不能为空")

    logger.info(f"收到总结请求，周期：{req.period}")
    return pipeline.run(raw_notes=req.raw_notes, period=req.period, author_name=req.author_name)


@app.post("/api/diagnose", response_model=DiagnoseResponse)
def diagnose(req: DiagnoseRequest):
    """
    跨周期诊断接口（V1.2新增）

    分析用户最近N条历史报告，识别反复出现的问题模式，
    并给出具体、可追溯依据的改进建议。
    """
    if not DEEPSEEK_API_KEY:
        raise HTTPException(status_code=400, detail="未配置 API Key，请检查 .env 文件")

    if req.report_ids:
        logger.info(f"收到诊断请求，用户手动勾选了 {len(req.report_ids)} 条报告")
    else:
        logger.info(f"收到诊断请求，分析最近 {req.report_count} 条报告")

    return diagnostic_pipeline.run(
        report_ids=req.report_ids,
        report_count=req.report_count,
        author_name=req.author_name
    )


@app.get("/api/reports", response_model=ReportListResponse)
def list_reports(
    page: int = Query(default=1, ge=1, description="页码，从1开始"),
    page_size: int = Query(default=10, ge=1, le=50, description="每页条数")
):
    """
    获取历史报告列表（分页），按创建时间倒序排列
    """
    all_reports = load_reports()
    all_reports.sort(key=lambda x: x.get("created_at", ""), reverse=True)

    total = len(all_reports)
    start = (page - 1) * page_size
    end = start + page_size
    page_reports = all_reports[start:end]

    summaries = []
    for r in page_reports:
        report_text = r.get("final_report", "")
        preview = report_text[:100].replace("\n", " ") + ("..." if len(report_text) > 100 else "")
        summaries.append(ReportSummary(
            report_id=r.get("report_id", ""),
            period=r.get("period", ""),
            created_at=r.get("created_at", ""),
            author_name=r.get("author_name", ""),
            preview=preview
        ))

    return ReportListResponse(
        reports=summaries,
        total=total,
        page=page,
        page_size=page_size
    )


@app.get("/api/reports/{report_id}")
def get_report(report_id: str):
    """获取单条历史报告的完整内容"""
    if report_id in report_cache:
        return report_cache[report_id]

    report = get_report_by_id(report_id)
    if not report:
        raise HTTPException(status_code=404, detail=f"报告 {report_id} 不存在")

    return report


@app.delete("/api/reports/{report_id}")
def delete_report(report_id: str):
    """删除指定报告"""
    report_cache.pop(report_id, None)

    success = delete_report_by_id(report_id)
    if not success:
        raise HTTPException(status_code=404, detail=f"报告 {report_id} 不存在")

    return {"message": f"报告 {report_id} 已删除", "success": True}


@app.get("/api/stats", response_model=StatsResponse)
def get_stats():
    """获取统计数据：总报告数、各周期分布、最近报告日期、本月报告数"""
    all_reports = load_reports()
    total = len(all_reports)

    period_dist = {}
    for r in all_reports:
        p = r.get("period", "未知")
        period_dist[p] = period_dist.get(p, 0) + 1

    latest_date = ""
    if all_reports:
        sorted_reports = sorted(all_reports, key=lambda x: x.get("created_at", ""), reverse=True)
        latest_raw = sorted_reports[0].get("created_at", "")
        if latest_raw:
            try:
                dt = datetime.fromisoformat(latest_raw)
                latest_date = dt.strftime("%Y年%m月%d日 %H:%M")
            except Exception:
                latest_date = latest_raw

    current_month = datetime.now().strftime("%Y-%m")
    this_month_count = sum(
        1 for r in all_reports
        if r.get("created_at", "").startswith(current_month)
    )

    return StatsResponse(
        total_reports=total,
        period_distribution=period_dist,
        latest_report_date=latest_date,
        this_month_count=this_month_count
    )


@app.post("/api/export/word")
def export_word(req: ExportRequest):
    """导出Word文档：根据 report_id 生成 .docx 文件返回给浏览器下载"""
    cached = report_cache.get(req.report_id) or get_report_by_id(req.report_id)

    if not cached:
        raise HTTPException(status_code=404, detail=f"报告 {req.report_id} 不存在")

    report_text = cached.get("final_report", "")
    period = cached.get("period", req.period)
    author_name = req.author_name or cached.get("author_name", "")

    logger.info(f"开始导出Word文档，报告ID：{req.report_id}")

    doc = DocxDocument()

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(f"{period}进度报告")
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(0x2B, 0x26, 0x20)

    meta_para = doc.add_paragraph()
    meta_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author_text = author_name if author_name else "个人"
    meta_run = meta_para.add_run(
        f"作者：{author_text}　　生成日期：{datetime.now().strftime('%Y年%m月%d日')}"
    )
    meta_run.font.size = Pt(10)
    meta_run.font.color.rgb = RGBColor(0x6B, 0x62, 0x56)

    doc.add_paragraph()

    for line in report_text.split("\n"):
        line = line.strip()
        if not line:
            doc.add_paragraph()
            continue
        if line.startswith("# "):
            h = doc.add_heading(line[2:], level=1)
            if h.runs:
                h.runs[0].font.color.rgb = RGBColor(0x2B, 0x26, 0x20)
        elif line.startswith("## "):
            h = doc.add_heading(line[3:], level=2)
            if h.runs:
                h.runs[0].font.color.rgb = RGBColor(0x8A, 0x5A, 0x44)
        elif line.startswith("---"):
            doc.add_paragraph()
        else:
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.font.size = Pt(11)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    filename = f"report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    logger.info(f"Word文档生成完成：{filename}")

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# 挂载静态文件目录
app.mount("/static", StaticFiles(directory="static"), name="static")
